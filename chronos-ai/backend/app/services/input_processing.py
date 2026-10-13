"""Input + document processing (spec sections 4.2 & 19).

Normalizes text, parses supported documents, enforces size limits, validates file
types by content where feasible, and treats all content as untrusted (never
executes macros/scripts, never fetches URLs).
"""
from __future__ import annotations

import email
from email import policy
from io import BytesIO

from ..core.config import settings


class InputError(ValueError):
    pass


def normalize_text(text: str) -> str:
    # Normalize line endings; preserve paragraph structure.
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # Collapse >2 blank lines to exactly 2 (keep paragraph boundaries).
    lines = text.split("\n")
    out: list[str] = []
    blanks = 0
    for ln in lines:
        if ln.strip() == "":
            blanks += 1
            if blanks <= 2:
                out.append("")
        else:
            blanks = 0
            out.append(ln.rstrip())
    normalized = "\n".join(out).strip()
    if len(normalized) > settings.MAX_TEXT_CHARACTERS:
        raise InputError(f"Input exceeds {settings.MAX_TEXT_CHARACTERS} characters.")
    return normalized


def process_text(text: str, email_subject: str = "", sender: str = "") -> tuple[str, bool]:
    """Return (normalized_text, truncated)."""
    if not text or not text.strip():
        raise InputError("Input text is empty.")
    normalized = normalize_text(text)
    header = ""
    if email_subject:
        header += f"Subject: {email_subject}\n"
    if sender:
        header += f"From: {sender}\n"
    combined = (header + "\n" + normalized) if header else normalized
    return combined, False


# --- document extraction -----------------------------------------------------

def extract_txt(content: bytes) -> str:
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError:
        return content.decode("latin-1", errors="replace")


def extract_eml(content: bytes) -> str:
    msg = email.message_from_bytes(content, policy=policy.default)
    subject = msg.get("subject", "")
    sender = msg.get("from", "")
    body_parts: list[str] = []
    if msg.is_multipart():
        for part in msg.walk():
            if part.get_content_type() == "text/plain":
                try:
                    body_parts.append(part.get_content())
                except Exception:
                    payload = part.get_payload(decode=True)
                    if payload:
                        body_parts.append(payload.decode("utf-8", errors="replace"))
    else:
        try:
            body_parts.append(msg.get_content())
        except Exception:
            payload = msg.get_payload(decode=True)
            if payload:
                body_parts.append(payload.decode("utf-8", errors="replace"))
    body = "\n".join(body_parts)
    return f"Subject: {subject}\nFrom: {sender}\n\n{body}".strip()


def extract_pdf(content: bytes) -> str:
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:  # pragma: no cover
        raise InputError("PDF support requires PyMuPDF (pymupdf).") from exc
    text_parts: list[str] = []
    with fitz.open(stream=content, filetype="pdf") as doc:
        for page in doc:
            text_parts.append(page.get_text("text"))
    return "\n".join(text_parts)


def extract_docx(content: bytes) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover
        raise InputError("DOCX support requires python-docx.") from exc
    document = docx.Document(BytesIO(content))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


_MAGIC = {
    b"%PDF-": "pdf",
    b"PK\x03\x04": "docx",  # docx is a zip
}


def sniff_type(filename: str, content: bytes) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    for magic, kind in _MAGIC.items():
        if content.startswith(magic):
            # zip could be docx or other; trust extension to disambiguate
            if kind == "docx" and ext not in ("docx",):
                continue
            return kind
    if ext in ("txt", "eml", "pdf", "docx"):
        return ext
    return ext or "unknown"


def process_file(filename: str, content: bytes, declared_type: str = "") -> tuple[str, bool]:
    if len(content) > settings.MAX_UPLOAD_BYTES:
        raise InputError(f"File exceeds {settings.MAX_UPLOAD_BYTES} bytes.")
    if not content:
        raise InputError("Uploaded file is empty.")

    kind = sniff_type(filename, content)
    try:
        if kind == "pdf":
            text = extract_pdf(content)
        elif kind == "docx":
            text = extract_docx(content)
        elif kind == "eml":
            text = extract_eml(content)
        elif kind == "txt":
            text = extract_txt(content)
        else:
            raise InputError(f"Unsupported or unrecognized file type: {kind}")
    except InputError:
        raise
    except Exception as exc:  # noqa: BLE001 — corrupt/invalid file
        raise InputError(f"Could not read {kind} file (it may be corrupt).") from exc

    if not text.strip():
        raise InputError("No readable text found in the document.")
    return normalize_text(text), False
